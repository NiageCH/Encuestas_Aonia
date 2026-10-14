"""Genera la plantilla Excel de asignación de profesores para el colegio (§7.2):
etl/plantillas/asignacion_docente_2526.xlsx

Una fila por docente, grupo y asignatura. Si dos docentes comparten asignatura y grupo
(codocencia o desdoble), van en dos filas. Hojas:
  * Asignación: la que rellena el colegio.
  * Instrucciones: qué poner en cada columna y un ejemplo.
  * Listas: grupos y asignaturas tal como aparecen en las encuestas 2025/26, para que los
    nombres casen. El desplegable avisa si se escribe otro valor, pero lo admite.

Uso (con la carga hecha en la base de datos local):
    .venv/bin/python etl/crear_plantilla_asignacion.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "etl"))
from cargar_2526 import url_bd  # noqa: E402

CURSO = "2025/26"
COLUMNAS = [
    ("Curso", 10, "Curso escolar, por ejemplo 2025/26."),
    ("Grupo", 10, "Grupo de alumnos tal como aparece en la hoja Listas, por ejemplo «Y9 A»."),
    ("Asignatura", 46, "Asignatura tal como aparece en la hoja Listas. Si el nombre del colegio es otro, escríbalo: se revisará a mano."),
    ("Apellidos", 24, "Apellidos del docente (se usan para el nombre del informe: Informe_2526_Apellido_Nombre.pdf)."),
    ("Nombre", 18, "Nombre del docente."),
    ("Correo", 30, "Correo del colegio del docente. Identifica a la persona: use siempre el mismo."),
    ("Departamento", 30, "Departamento al que pertenece el docente (para los informes de jefes de departamento)."),
]
# Colores de la intranet (:root): navy y soft.
NAVY, SOFT = "1A2C4A", "EAF0F6"


def listas() -> tuple[list[str], list[str]]:
    import psycopg
    with psycopg.connect(url_bd(False)) as conn:
        grupos = [r[0] for r in conn.execute(
            """select g.nombre from public.grupo g join public.curso c on c.id = g.curso_id
               where c.nombre = %s order by g.etapa, length(g.nivel), g.nivel, g.nombre""", (CURSO,))]
        asignaturas = [r[0] for r in conn.execute(
            """select a.nombre from public.asignatura a join public.curso c on c.id = a.curso_id
               where c.nombre = %s order by a.nombre""", (CURSO,))]
    return grupos, asignaturas


def main() -> None:
    grupos, asignaturas = listas()
    libro = Workbook()
    cab_fuente = Font(bold=True, color="FFFFFF")
    cab_relleno = PatternFill("solid", fgColor=NAVY)

    hoja = libro.active
    hoja.title = "Asignación"
    for i, (titulo, ancho, _) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=i, value=titulo)
        celda.font, celda.fill = cab_fuente, cab_relleno
        celda.alignment = Alignment(vertical="center")
        hoja.column_dimensions[celda.column_letter].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:{hoja.cell(row=1, column=len(COLUMNAS)).column_letter}1"

    ls = libro.create_sheet("Listas")
    ls["A1"], ls["B1"] = "Grupos 2025/26", "Asignaturas 2025/26 (nombres de las encuestas)"
    for c in ("A1", "B1"):
        ls[c].font, ls[c].fill = cab_fuente, cab_relleno
    for i, g in enumerate(grupos, start=2):
        ls.cell(row=i, column=1, value=g)
    for i, a in enumerate(asignaturas, start=2):
        ls.cell(row=i, column=2, value=a)
    ls.column_dimensions["A"].width, ls.column_dimensions["B"].width = 18, 70

    def desplegable(rango_lista: str, columna: str, nombre: str) -> None:
        dv = DataValidation(type="list", formula1=rango_lista, allow_blank=True, showErrorMessage=True,
                            errorStyle="warning", errorTitle=f"{nombre} no encontrado",
                            error=f"Este {nombre.lower()} no está en la hoja Listas. Se puede dejar, pero se revisará a mano.")
        hoja.add_data_validation(dv)
        dv.add(f"{columna}2:{columna}2000")

    desplegable(f"=Listas!$A$2:$A${len(grupos) + 1}", "B", "Grupo")
    desplegable(f"=Listas!$B$2:$B${len(asignaturas) + 1}", "C", "Asignatura")
    dv_curso = DataValidation(type="list", formula1=f'"{CURSO}"', allow_blank=True)
    hoja.add_data_validation(dv_curso)
    dv_curso.add("A2:A2000")

    ins = libro.create_sheet("Instrucciones", 0)
    ins.column_dimensions["A"].width, ins.column_dimensions["B"].width = 16, 100
    filas = [
        ("Asignación de profesores por grupo y asignatura · curso 2025/26", ""),
        ("", ""),
        ("Para qué sirve", "Con esta tabla la plataforma de encuestas sabe qué docente imparte cada asignatura en cada grupo, "
                           "para preparar el informe de cada profesor con las encuestas de sus alumnos."),
        ("Cómo rellenarla", "En la hoja «Asignación», una fila por docente, grupo y asignatura. Si dos docentes comparten "
                            "asignatura y grupo (codocencia o desdoble), escriba dos filas."),
        ("", "Los nombres de grupo y asignatura conviene copiarlos de la hoja «Listas»: son los que aparecen en las encuestas. "
             "Si escribe otro, se acepta, pero se revisará a mano."),
        ("", "No añada ni quite columnas. No hace falta ordenar las filas."),
        ("", ""),
        ("Columna", "Qué poner"),
        *[(t, d) for t, _, d in COLUMNAS],
        ("", ""),
        ("Ejemplo", "2025/26 · Y9 A · Matemáticas · García López · Ana · ana.garcia@sfpaula.com · Matemáticas"),
        ("", "(Ejemplo inventado: no lo copie en la hoja Asignación.)"),
        ("", ""),
        ("Privacidad", "Esta tabla solo se usa para atribuir los resultados agregados de las encuestas. Los alumnos no "
                       "aparecen en ella y ningún resultado con menos de 5 respuestas se muestra en los informes."),
    ]
    for i, (a, b) in enumerate(filas, start=1):
        ins.cell(row=i, column=1, value=a).font = Font(bold=True, color=NAVY)
        ins.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
    ins["A1"].font = Font(bold=True, size=14, color=NAVY)
    for c in ("A8", "B8"):
        ins[c].font, ins[c].fill = cab_fuente, cab_relleno
    for i in range(9, 9 + len(COLUMNAS)):
        ins.cell(row=i, column=1).fill = PatternFill("solid", fgColor=SOFT)

    destino = RAIZ / "etl" / "plantillas" / "asignacion_docente_2526.xlsx"
    libro.active = libro.sheetnames.index("Asignación")
    libro.save(destino)
    print(f"{destino.relative_to(RAIZ)}: {len(grupos)} grupos y {len(asignaturas)} asignaturas en la hoja Listas")


if __name__ == "__main__":
    main()
