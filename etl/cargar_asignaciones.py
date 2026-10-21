"""Carga la asignación de profesores por grupo y asignatura (§7.2) desde el Excel del colegio
(plantilla: etl/plantillas/asignacion_docente_2526.xlsx).

Uso:
    .venv/bin/python etl/cargar_asignaciones.py RUTA_DEL_EXCEL [--ensayo]

- Cada docente se identifica por su correo (si cambia el nombre, se actualiza).
- Grupos y asignaturas se casan con los de las encuestas del curso (sin distinguir
  mayúsculas, acentos ni espacios). Lo que no casa NO se carga: va al informe
  datos_locales/informes/asignaciones_revision.md para revisarlo a mano.
- Repetible: sustituye todas las asignaciones del curso.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "etl"))
from cargar_2526 import url_bd  # noqa: E402
from encuesta_facil import normalizar_grupo, sin_acentos  # noqa: E402
from departamentos_2526 import OTROS  # noqa: E402

COLUMNAS = ["Curso", "Grupo", "Asignatura", "Apellidos", "Nombre", "Correo", "Departamento"]


def clave(s: str) -> str:
    return re.sub(r"\s+", " ", sin_acentos(str(s or "")).strip().lower())


def leer(ruta: Path) -> list[dict]:
    import openpyxl
    hoja = openpyxl.load_workbook(ruta, read_only=True, data_only=True)["Asignación"]
    filas = list(hoja.iter_rows(values_only=True))
    cab = [str(c or "").strip() for c in filas[0]]
    if cab[:len(COLUMNAS)] != COLUMNAS:
        raise ValueError(f"Columnas inesperadas: {cab}. Deben ser {COLUMNAS}")
    # Hoja opcional «Departamentos» (la genera convertir_alexia_2526.py): manda sobre la de
    # asignación en departamento, apellidos y nombre de cada correo.
    por_correo = {}
    if "Departamentos" in hoja.parent.sheetnames:
        deps = list(hoja.parent["Departamentos"].iter_rows(values_only=True))
        cd = [str(c or "").strip() for c in deps[0]]
        for f in deps[1:]:
            r = {c: (str(v).strip() if v is not None else "") for c, v in zip(cd, f)}
            if r.get("Correo"):
                por_correo[r["Correo"].lower()] = r
    salida = []
    for n, f in enumerate(filas[1:], start=2):
        d = {c: (str(v).strip() if v is not None else "") for c, v in zip(COLUMNAS, f)}
        if any(d.values()):
            for c in ("Departamento", "Apellidos", "Nombre"):
                d[c] = por_correo.get(d["Correo"].lower(), {}).get(c) or d[c]
            d["Situación"] = por_correo.get(d["Correo"].lower(), {}).get("Situación") or "Alta"
            d["fila"] = n
            salida.append(d)
    return salida


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel", type=Path)
    ap.add_argument("--ensayo", action="store_true", help="valida y muestra el informe, sin cargar")
    ap.add_argument("--remoto", action="store_true")
    args = ap.parse_args()

    import psycopg
    filas = leer(args.excel)
    problemas: list[str] = []
    validas = []
    with psycopg.connect(url_bd(args.remoto)) as conn, conn.cursor() as cur:
        cur.execute("select nombre, id from public.curso")
        cursos = dict(cur.fetchall())
        for f in filas:
            falta = [c for c in COLUMNAS if not f[c]]
            if falta:
                problemas.append(f"Fila {f['fila']}: faltan {', '.join(falta)}")
                continue
            if f["Curso"] not in cursos:
                problemas.append(f"Fila {f['fila']}: curso {f['Curso']!r} no cargado")
                continue
            if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", f["Correo"]):
                problemas.append(f"Fila {f['fila']}: correo no válido {f['Correo']!r}")
                continue
            validas.append(f)

        por_curso = Counter(f["Curso"] for f in validas)
        grupos, asignaturas = {}, {}
        for curso in por_curso:
            cid = cursos[curso]
            cur.execute("select nombre, id from public.grupo where curso_id = %s", (cid,))
            grupos[curso] = {clave(n): i for n, i in cur.fetchall()}
            cur.execute("select nombre, id from public.asignatura where curso_id = %s", (cid,))
            asignaturas[curso] = {clave(n): (n, i) for n, i in cur.fetchall()}

        casadas = []
        sin_casar: Counter = Counter()
        for f in validas:
            try:
                g = grupos[f["Curso"]].get(clave(normalizar_grupo(f["Grupo"]) or ""))
            except ValueError:
                g = None
            a = asignaturas[f["Curso"]].get(clave(f["Asignatura"]))
            if g is None:
                problemas.append(f"Fila {f['fila']}: grupo {f['Grupo']!r} no está en las encuestas del curso")
            if a is None:
                sin_casar[f["Asignatura"]] += 1
            if g is not None and a is not None:
                casadas.append((f, g, a[1]))

        informe = [f"# Revisión de la asignación de profesores ({args.excel.name})", "",
                   f"- Filas leídas: {len(filas)}", f"- Filas cargables: {len(casadas)}",
                   f"- Filas con problemas: {len(filas) - len(casadas)}", ""]
        if sin_casar:
            informe += ["## Asignaturas que no casan con las encuestas", "",
                        "Hay que decir a qué asignatura de la hoja Listas corresponde cada una:", ""]
            informe += [f"- {a!r} ({n} filas)" for a, n in sin_casar.most_common()]
            informe.append("")
        if problemas:
            informe += ["## Otros problemas", ""] + [f"- {p}" for p in problemas] + [""]
        destino = RAIZ / "datos_locales" / "informes" / "asignaciones_revision.md"
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text("\n".join(informe))
        print("\n".join(informe[:5]))
        print(f"Informe: {destino.relative_to(RAIZ)}")
        if args.ensayo:
            return

        for curso in por_curso:
            cur.execute("delete from public.asignacion_docente where curso_id = %s", (cursos[curso],))
        for f, g, a in casadas:
            cid = cursos[f["Curso"]]
            cur.execute("""insert into public.departamento (curso_id, nombre, oficial) values (%s, %s, %s)
                           on conflict (curso_id, nombre) do update set oficial = excluded.oficial""",
                        (cid, f["Departamento"], f["Departamento"] != OTROS))
            cur.execute("select id from public.departamento where curso_id = %s and nombre = %s", (cid, f["Departamento"]))
            dep = cur.fetchone()[0]
            cur.execute(
                """insert into public.persona_responsable (nombre, apellidos, correo, departamento_id, situacion)
                   values (%s, %s, lower(%s), %s, %s)
                   on conflict (correo) do update set nombre = excluded.nombre, apellidos = excluded.apellidos,
                     departamento_id = excluded.departamento_id, situacion = excluded.situacion
                   returning id""", (f["Nombre"], f["Apellidos"], f["Correo"], dep,
                                     "baja" if f["Situación"].lower() == "baja" else "alta"))
            pid = cur.fetchone()[0]
            cur.execute("insert into public.asignacion_docente (curso_id, grupo_id, asignatura_id, docente_id) values (%s, %s, %s, %s) on conflict do nothing",
                        (cid, g, a, pid))
        conn.commit()
        print(f"Cargadas {len(casadas)} asignaciones.")


if __name__ == "__main__":
    main()
