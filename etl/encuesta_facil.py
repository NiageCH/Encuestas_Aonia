"""Lectura y normalización de las exportaciones de Encuesta Fácil (CSV o XLSX).

Formato de la exportación:
  * unas filas de metadatos;
  * la fila de preguntas, que empieza por "IDCuestionario" (se localiza por contenido);
  * la fila de subtítulos justo debajo: "Valor" (columna numérica de una valoración, seguida
    de la columna con su etiqueta, cuyo subtítulo es el texto del ítem), "Respuesta",
    "Respuesta abierta" o el texto de una opción;
  * los datos.

Este módulo no sabe nada de cuestionarios concretos: devuelve columnas tipadas. La
clasificación (qué es cada bloque) está en cuestionarios_2526.py.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

# Columnas de metadatos que se descartan siempre (identificadores, fechas, IP, correo,
# nombre, campos libres). "Abandonado" solo se usa para contar.
METADATOS = {
    "IDCuestionario", "Inicio", "Fin", "Abandonado", "Dirección IP", "EMail", "Nombre",
    "Apellido 1", "Apellido 2", "Contenido libre", "Parametro de usuario 1",
    "Parametro de usuario 2", "Parametro de usuario ID",
}


@dataclass
class Columna:
    """Una pregunta o ítem de la exportación."""
    tipo: str              # "valor" (con etiqueta al lado), "respuesta", "abierta" u "opcion"
    col: int               # columna del dato (en "valor", la del número)
    cabecera: str          # texto de la pregunta del bloque al que pertenece
    texto: str             # texto del ítem o de la opción ("" en "respuesta")
    col_etiqueta: Optional[int] = None


@dataclass
class Tabla:
    origen: str
    preguntas: List[str]
    subtitulos: List[str]
    filas: List[List[str]]
    col_abandonado: int
    columnas: List[Columna] = field(default_factory=list)

    def valor(self, fila: List[str], col: int) -> str:
        return fila[col].strip() if col < len(fila) and fila[col] is not None else ""

    def abandonada(self, fila: List[str]) -> bool:
        return bool(self.valor(fila, self.col_abandonado))


def _leer_filas(ruta: Path) -> List[List[str]]:
    if ruta.suffix.lower() == ".csv":
        with ruta.open(encoding="utf-8-sig", newline="") as f:
            return [list(r) for r in csv.reader(f)]
    if ruta.suffix.lower() == ".xlsx":
        import openpyxl
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        hoja = libro[libro.sheetnames[0]]
        filas = []
        for r in hoja.iter_rows(values_only=True):
            filas.append([_texto_celda(c) for c in r])
        return filas
    raise ValueError(f"Formato no admitido: {ruta.name}")


def _texto_celda(c) -> str:
    if c is None:
        return ""
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c)


def leer(ruta: Path) -> Tabla:
    filas = _leer_filas(ruta)
    try:
        i = next(k for k, r in enumerate(filas) if r and r[0].strip() == "IDCuestionario")
    except StopIteration:
        raise ValueError(f"{ruta.name}: no encuentro la fila de preguntas (IDCuestionario)")
    ancho = max(len(r) for r in filas)
    preguntas = [c.strip() for c in filas[i]] + [""] * (ancho - len(filas[i]))
    subtitulos = [c.strip() for c in filas[i + 1]] + [""] * (ancho - len(filas[i + 1]))
    datos = [r + [""] * (ancho - len(r)) for r in filas[i + 2:] if any(c.strip() for c in r)]
    tabla = Tabla(ruta.name, preguntas, subtitulos, datos, preguntas.index("Abandonado"))
    tabla.columnas = _columnas(tabla)
    return tabla


def _columnas(t: Tabla) -> List[Columna]:
    cols: List[Columna] = []
    cabecera = ""
    j = 0
    while j < len(t.preguntas):
        if t.preguntas[j]:
            cabecera = t.preguntas[j]
        if j < 13 or cabecera in METADATOS:
            j += 1
            continue
        sub = t.subtitulos[j]
        if sub == "Valor":
            cols.append(Columna("valor", j, cabecera, t.subtitulos[j + 1], col_etiqueta=j + 1))
            j += 2
            continue
        if sub == "Respuesta abierta":
            cols.append(Columna("abierta", j, cabecera, ""))
        elif sub == "Respuesta":
            cols.append(Columna("respuesta", j, cabecera, ""))
        else:
            cols.append(Columna("opcion", j, cabecera, sub))
        j += 1
    return cols


# Normalización ----------------------------------------------------------------------------

ERRATAS = {
    "trabajdo": "trabajado",
    "instrudio": "instruido",
}


def sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def normalizar_texto(s: str) -> str:
    """Texto comparable: sin puntos suspensivos iniciales, sin espacios repetidos, sin punto
    final y con erratas conocidas corregidas. Conserva mayúsculas y acentos."""
    s = s.replace(" ", " ").strip()
    s = re.sub(r"^(\.{2,}|…)\s*", "", s)
    s = re.sub(r"\s+", " ", s).strip().rstrip(".").strip()
    for mal, bien in ERRATAS.items():
        s = re.sub(rf"\b{mal}\b", bien, s)
    return s


def slug(s: str, largo: int = 60) -> str:
    base = sin_acentos(normalizar_texto(s)).lower()
    base = re.sub(r"[^a-z0-9]+", "_", base).strip("_")
    if len(base) > largo:
        base = base[:largo].rsplit("_", 1)[0]
    return base


def frase(s: str) -> str:
    """'LENGUA ESPAÑOLA o LENGUA ADAPTADA' -> 'Lengua española o lengua adaptada'."""
    s = re.sub(r"\s+", " ", s).strip()
    return s[:1].upper() + s[1:].lower()


def normalizar_grupo(valor: str, nivel_por_defecto: Optional[str] = None) -> Optional[str]:
    """'Y8A' -> 'Y8 A'; 'A' con nivel 'Y10' -> 'Y10 A'; '' -> None."""
    v = valor.strip().upper().replace("-", " ")
    if not v:
        return None
    m = re.fullmatch(r"(Y\d{1,2})\s*([A-Z])", v)
    if m:
        return f"{m.group(1)} {m.group(2)}"
    if re.fullmatch(r"[A-Z]", v) and nivel_por_defecto:
        return f"{nivel_por_defecto} {v}"
    raise ValueError(f"Grupo no reconocido: {valor!r}")


def etapa_de_nivel(nivel: str) -> str:
    n = int(nivel[1:])
    if 1 <= n <= 5:
        return "primaria"
    if 6 <= n <= 10:
        return "pai"
    if 11 <= n <= 12:
        return "bachillerato"
    raise ValueError(f"Nivel sin etapa: {nivel}")
