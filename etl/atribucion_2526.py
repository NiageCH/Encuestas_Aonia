"""Atribución de respuestas de alumnos a un docente (decisión del 21/10/2026, §7.2).

Algunos cuestionarios preguntan «nombre de tu profesor de…» y el de Y6 a Y9 «nombre de la
optativa». Con el listado de Alexia, el texto se casa con un docente de ese grupo y esa
asignatura (o con la optativa, si tiene un solo docente). Solo se guarda el docente
resultante; el texto nunca se carga ni se escribe en ningún informe.

Se aplica solo:
- en las clases con más de un docente (en las demás no cambia nada);
- en las columnas en las que el nombre casa con un único docente en el 90 % o más de los
  alumnos que la rellenaron;
- no en Inglés de Y10 y Y12, hasta que el colegio aclare el posible hueco de Alexia.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

from convertir_alexia_2526 import CORRESPONDENCIA, grupo as grupo_alexia, k, partir_nombre
from encuesta_facil import sin_acentos

ALEXIA = Path(__file__).resolve().parent.parent / "datos_locales" / "fuentes" / "Grupos_asignaturas_docentes_2025-26.xlsx"
UMBRAL = 0.90

# (cuestionario, columna «profesor de …» sin acentos) → asignatura de la encuesta.
COLUMNAS = {
    ("ALU_Y10", "lengua espanola"): "Lengua española",
    ("ALU_Y10", "ingles"): "Inglés o adquisición de lengua inglesa",
    ("ALU_Y10", "frances"): "Francés",
    ("ALU_Y11", "ingles"): "Inglés",
    ("ALU_Y11", "historia"): "Historia de américa o historia nm",
    ("ALU_Y11", "matematicas"): "Matemáticas i o matemáticas apl. cc.ss. i o matemáticas: análisis y enfoques nm o matemáticas: análisis y enfoques ns",
    ("ALU_Y11", "quimica"): "Química",
    ("ALU_Y12", "ingles"): "Inglés o inglés a ns o inglés b ns",
    ("ALU_Y12", "matematicas"): "Matemáticas ii o matemáticas apl. cc.ss. ii o matemáticas: análisis y enfoques nm o matemáticas: análisis y enfoques ns",
    ("ALU_Y12", "quimica"): "Química o química nm o química ns",
    ("ALU_Y12", "biologia"): "Biología o biología nm o biología ns",
    ("ALU_Y6Y9", "optativa"): "Optativa",
}

EXCLUIDAS = {
    ("ALU_Y10", "ingles"): "posible hueco en Alexia, pendiente del colegio",
    ("ALU_Y12", "ingles"): "posible hueco en Alexia, pendiente del colegio",
}


def palabras(s: str) -> Set[str]:
    return set(re.sub(r"[^a-z ]", " ", sin_acentos(str(s or "")).lower()).split())


@dataclass
class Alexia:
    nombre: Dict[str, str] = field(default_factory=dict)                         # correo → nombre completo
    por_grupo: Dict[str, Dict[str, Set[str]]] = field(default_factory=lambda: defaultdict(lambda: defaultdict(set)))
    optativas: Dict[str, Set[str]] = field(default_factory=lambda: defaultdict(set))   # optativa → correos

    def tokens(self, correo: str) -> Set[str]:
        return {t for t in palabras(self.nombre[correo]) if len(t) >= 3 and t not in ("del", "las", "los")}


def leer_alexia(ruta: Path = ALEXIA) -> Alexia:
    import openpyxl
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    a = Alexia()
    filas = list(libro["Detalle docente"].iter_rows(values_only=True))
    cab = [str(c or "").strip() for c in filas[0]]
    i = {c: cab.index(c) for c in ["Grupo", "Asignatura", "Docente", "Correo"]}
    for f in filas[1:]:
        correo = str(f[i["Correo"]] or "").strip().lower()
        g = grupo_alexia(f[i["Grupo"]])
        if correo:
            a.nombre[correo] = str(f[i["Docente"]]).strip()
            if g:
                a.por_grupo[g][correo].add(str(f[i["Asignatura"]]))
    filas = list(libro["Optativas y extraescolares"].iter_rows(values_only=True))
    cab = [str(c or "").strip() for c in filas[0]]
    i = {c: cab.index(c) for c in ["Programa", "Asignatura", "Docente", "Correo"]}
    for f in filas[1:]:
        if str(f[i["Programa"]]).strip() == "Optativas Y6-Y9" and f[i["Correo"]]:
            correo = str(f[i["Correo"]]).strip().lower()
            a.nombre.setdefault(correo, str(f[i["Docente"]]).strip())
            a.optativas[str(f[i["Asignatura"]]).strip()].add(correo)
    return a


@dataclass
class Resumen:
    cuestionario: str
    columna: str
    alumnos: int = 0
    rellenan: int = 0
    casan: int = 0
    atribuidas: int = 0
    aplicada: bool = False
    motivo: str = ""


def candidatos(a: Alexia, cuestionario: str, g: str, destino: str) -> List[str]:
    tabla = CORRESPONDENCIA.get(cuestionario, {})
    return sorted(c for c, asigs in a.por_grupo.get(g, {}).items() if any(tabla.get(k(x)) == destino for x in asigs))


def atribuir(a: Alexia, cuestionario: str, respuestas) -> Tuple[List[Tuple[int, str, str]], List[Resumen]]:
    """Devuelve [(índice de la respuesta, asignatura de la encuesta, correo)] y el resumen por columna."""
    salida: List[Tuple[int, str, str]] = []
    resumenes: List[Resumen] = []
    claves_opt = {o: {w for w in palabras(o) if len(w) >= 4} for o in a.optativas}
    for (q, col), destino in COLUMNAS.items():
        if q != cuestionario:
            continue
        res = Resumen(q, col)
        propuestas: List[Tuple[int, str, str]] = []
        for n, r in enumerate(respuestas):
            if col == "optativa":
                res.alumnos += 1
            elif not r.grupo:
                continue
            else:
                res.alumnos += 1
            texto = r.textos.get(col)
            if not texto or not palabras(texto):
                continue
            res.rellenan += 1
            pal = palabras(texto)
            if col == "optativa":
                hits = [o for o, ks in claves_opt.items() if ks & pal]
                if len(hits) == 1:
                    res.casan += 1
                    if len(a.optativas[hits[0]]) == 1:
                        propuestas.append((n, destino, next(iter(a.optativas[hits[0]]))))
                continue
            cand = candidatos(a, q, r.grupo, destino)
            hits = [c for c in cand if a.tokens(c) & pal]
            if len(hits) == 1:
                res.casan += 1
                if len(cand) > 1:
                    propuestas.append((n, destino, hits[0]))
        tasa = res.casan / res.rellenan if res.rellenan else 0
        if (q, col) in EXCLUIDAS:
            res.motivo = EXCLUIDAS[(q, col)]
        elif tasa < UMBRAL:
            res.motivo = f"casa en menos del {UMBRAL:.0%}"
        else:
            res.aplicada = True
            res.atribuidas = len(propuestas)
            salida += propuestas
        resumenes.append(res)
    return salida, resumenes


def nombre_y_apellidos(a: Alexia, correo: str) -> Tuple[str, str]:
    nombre, apellidos, _ = partir_nombre(a.nombre[correo], correo)
    return nombre, apellidos
