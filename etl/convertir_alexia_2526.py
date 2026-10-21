"""Convierte el listado de Alexia (grupos, asignaturas y docentes 2025/26) al formato de la
plantilla de asignación (etl/plantillas/asignacion_docente_2526.xlsx), listo para
etl/cargar_asignaciones.py.

Uso:
    .venv/bin/python etl/convertir_alexia_2526.py [RUTA_ALEXIA] [--salida RUTA]

Por defecto lee datos_locales/fuentes/Grupos_asignaturas_docentes_2025-26.xlsx (hoja
«Detalle docente») y escribe datos_locales/asignacion_2526_desde_alexia.xlsx con:
- Asignación: una fila por docente, grupo y asignatura de la encuesta.
- Departamentos: departamento PROPUESTO por docente (según lo que imparte), para que el
  colegio lo confirme o corrija en esa hoja. Al cargar, cargar_asignaciones.py toma de esta
  hoja el departamento, los apellidos y el nombre de cada correo.
- Codocencia: asignaturas de la encuesta con más de un docente en el mismo grupo.
- Sin encuesta: asignaturas de Alexia que no se preguntan a los alumnos (no se cargan).

Las encuestas agrupan varias materias de Alexia en una sola pregunta («Lengua española o
lengua adaptada»): la tabla CORRESPONDENCIA dice a qué pregunta va cada una, por cuestionario.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "etl"))
from cargar_2526 import url_bd  # noqa: E402
from encuesta_facil import sin_acentos  # noqa: E402
from departamentos_2526 import INGLES_Y_ADQUISICION, OFICIALES, OTROS, RENOMBRAR  # noqa: E402

CURSO = "2025/26"
ORIGEN = RAIZ / "datos_locales" / "fuentes" / "Grupos_asignaturas_docentes_2025-26.xlsx"
SALIDA = RAIZ / "datos_locales" / "asignacion_2526_desde_alexia.xlsx"
REVISION = RAIZ / "datos_locales" / "Revision_profesores_departamentos_2526_cumplimentado.xlsx"
REVISAR_INGLES = RAIZ / "datos_locales" / "revisar_ingles_adquisicion_2526.xlsx"

# Reparto del antiguo «Inglés y Adquisición de Lenguas» según lo que imparte cada docente.
INGLESAS = r"^(lengua inglesa y literatura|ingles a |ingles i+ hsd)"
ADQUISICION = r"adquisicion de lengua inglesa|ingles b |frances|aleman|chino|arabe|inmersion|adquisicion de lengua espanola|espanol b"


def k(s: str) -> str:
    """Clave de comparación: sin acentos, minúsculas, sin sufijo de nivel («Matemáticas-Y7»)."""
    s = re.sub(r"-\s*Y\d{1,2}$", "", str(s or "").strip())
    return re.sub(r"\s+", " ", sin_acentos(s).lower()).strip()


def tabla(pares: dict[str, list[str]]) -> dict[str, str]:
    return {k(a): encuesta for encuesta, alexia in pares.items() for a in alexia}


# Asignatura de Alexia → asignatura de la encuesta, por cuestionario de alumnos.
CORRESPONDENCIA = {
    "ALU_Y2Y5": tabla({
        "Lengua española": ["Lengua Española y Literatura"],
        "Español": ["Adquisición de Lengua Española"],
        "Inglés": ["Lengua Inglesa y Literatura"],
        "Matemáticas": ["Matemáticas"],
        "Ciencias": ["Ciencias Naturales y Sociales"],
        "Artes visuales": ["Artes Visuales"],
        "Chino": ["Chino"], "Francés": ["Francés"], "Árabe": ["Árabe"],
        "Dramatización": ["Dramatización"],
        "Educación física": ["Educación Personal, Social y Física"],
        "Exposición pep": ["Exposición"],
        "Música (grupo completo)": ["Música"],
        "Steam": ["STEaM"],
        "Ajedrez o conjunto instrumental o mateolímpicos o teatro": ["Ajedrez", "Conjunto Instrumental", "Mateolímpicos", "Teatro"],
    }),
    "ALU_Y6Y9": tabla({
        "Lengua española o lengua adaptada": ["Lengua Española y Literatura", "Lengua Española y Literatura- Adaptada"],
        "Español": ["Adquisición de Lengua Española"],
        "Inglés o adquisición de inglés": ["Lengua Inglesa y Literatura", "Adquisición de Lengua Inglesa"],
        "Matemáticas": ["Matemáticas"], "Ciencias": ["Ciencias"], "Diseño": ["Diseño"],
        "Alemán": ["Alemán"], "Chino": ["Chino"], "Francés": ["Francés"],
        "Artes visuales": ["Artes Visuales"],
        "Individuos y sociedades": ["Individuos y Sociedades (Ciencias Sociales)"],
        "Educación física": ["Educación Física y de la Salud", "Educación Física y para la Salud"],
        "Música o música pim": ["Música", "Música PIM"],
    }),
    "ALU_Y6Y10": tabla({
        "Lengua española y literatura o lengua adaptada": ["Lengua Española y Literatura", "Lengua Española y Literatura- Adaptada"],
        "Adquisición de inglés o lengua inglesa y literatura": ["Lengua Inglesa y Literatura", "Adquisición de Lengua Inglesa"],
        "Artes visuales o música o pim": ["Artes Visuales", "Música", "Música PIM"],
        "Ciencias": ["Ciencias"], "Diseño": ["Diseño"], "Matemáticas": ["Matemáticas"],
        "Educación física en español o en inglés": ["Educación Física y de la Salud", "Educación Física y para la Salud"],
        "Francés o alemán o chino o árabe o español b": ["Francés", "Alemán", "Chino", "Árabe", "Adquisición de Lengua Española"],
        "Individuos y sociedades": ["Individuos y Sociedades (Ciencias Sociales)"],
        "Planificación y reflexión": ["Planificación y Reflexión"],
        "Programación": ["Programación", "Programacion"],
        "Proyecto personal": ["Proyecto Personal"],
    }),
    "ALU_Y10": tabla({
        "Lengua española": ["Lengua Española y Literatura", "Lengua Española y Literatura- Adaptada"],
        "Lengua española b": ["Adquisición de Lengua Española"],
        "Inglés o adquisición de lengua inglesa": ["Lengua Inglesa y Literatura", "Adquisición de Lengua Inglesa"],
        "Alemán": ["Alemán"], "Chino": ["Chino"], "Francés": ["Francés"],
        "Ciencias": ["Ciencias"], "Diseño": ["Diseño"], "Matemáticas": ["Matemáticas"],
        "Música": ["Música"], "Artes visuales": ["Artes Visuales"],
        "Educación física": ["Educación Física y de la Salud", "Educación Física y para la Salud"],
        "Individuos y sociedades": ["Individuos y Sociedades (Ciencias Sociales)"],
    }),
    "ALU_Y11": tabla({
        "Anatomía aplicada": ["Anatomía Aplicada I"],
        "Biología i o biología ns": ["Biología I HSD", "Biología NS"],
        "Cas": ["Creatividad, Actividad y Servicio"],
        "Derecho y relaciones internacionales": ["Derecho y Relaciones Internacionales"],
        "Dibujo técnico i": ["Dibujo Técnico I"],
        "Diseño asistido por ordenador (cad)": ["Diseño Asistido por Ordenador (CAD)"],
        "Economía": ["Economía"],
        "Física i o física ns": ["Física I HSD", "Física NS"],
        "Francés": ["Francés I"],
        "Historia de américa o historia nm": ["Historia de América", "Historia NM"],
        "Inglés": ["Inglés I HSD", "Inglés A Lengua y Literatura NS", "Inglés B Adquisición NS"],
        "Lengua española y literatura o español a literatura nm o español b": [
            "Lengua Española y Literatura", "Español A Literatura NM", "Español B Adquisición NM"],
        "Matemáticas i o matemáticas apl. cc.ss. i o matemáticas: análisis y enfoques nm o matemáticas: análisis y enfoques ns": [
            "Matemáticas I HSD", "Matemáticas Aplicadas a las CC.SS. I",
            "Matemáticas: Análisis y Enfoques NM", "Matemáticas: Análisis y Enfoques NS"],
        "Química": ["Química I HSD", "Química NM", "Química NS"],
        "Teoría del conocimiento": ["Teoría del Conocimiento"],
    }),
    "ALU_Y12": tabla({
        "Ampliación de matemáticas aplicadas a las cc.ss.": ["Ampliación MatAplCCSS"],
        "Biología o biología nm o biología ns": ["Biología II HSD", "Biología NM", "Biología NS"],
        "Dibujo artístico ii": ["Dibujo Artístico II"],
        "Dibujo técnico ii": ["Dibujo Técnico II"],
        "Diseño asistido por ordenador (cad)": ["Diseño Asistido por Ordenador (CAD)"],
        "Empresa y diseño de modelos de negocio o gestión empresarial nm": [
            "Empresa y Diseño de Modelos de Negocio", "Gestión Empresarial NM"],
        "Español b ns": ["Español B Adquisición NS"],
        "Física o física ns": ["Física II HSD", "Física NS"],
        "Francés": ["Francés II"],
        "Fundamentos artísticos ii": ["Fundamentos Artísticos II"],
        "Geografía ii": ["Geografía II HSD"],
        "Historia del arte": ["Historia del Arte II"],
        "Inglés o inglés a ns o inglés b ns": ["Inglés II HSD", "Inglés A Lengua y Literatura NS", "Inglés B Adquisición NS"],
        "Lengua española y literatura o literatura nm o literatura ns": [
            "Lengua Española y Literatura HSD", "Español A Literatura NM", "Español A Literatura NS"],
        "Matemáticas ii o matemáticas apl. cc.ss. ii o matemáticas: análisis y enfoques nm o matemáticas: análisis y enfoques ns": [
            "Matemáticas II HSD", "Matemáticas Aplicadas a las Ciencias Sociales II",
            "Matemáticas: Análisis y Enfoques NM", "Matemáticas: Análisis y Enfoques NS"],
        "Política global ns": ["Política global NS"],
        "Química o química nm o química ns": ["Química II HSD", "Química NM", "Química NS"],
        "Teoría del conocimiento": ["Teoría del Conocimiento"],
    }),
}

# Preguntas de la encuesta que no se pueden atribuir: sus docentes no van por grupo de clase.
NO_ATRIBUIBLES = {
    "Inmersión lingüística (a1 o a2 o b1 o b2)": "Los grupos de inmersión (ALE A1 a B2) no son grupos de clase: no se sabe qué docente tuvo cada alumno.",
    "Optativa": "Las optativas de Y6 a Y9 mezclan alumnos de varios grupos y cada alumno valora la suya.",
}

# Departamento propuesto según la asignatura de Alexia (el colegio lo confirma).
DEPARTAMENTO = [
    ("Lengua Española y Literatura", r"lengua espanola|espanol|lel |suplemento lengua"),
    ("Inglés y Adquisición de Lenguas", r"ingles|lengua inglesa|frances|aleman|chino|arabe|inmersion"),
    ("Matemáticas", r"matematic|mat ?apl|matapl|\bmae\b|mateolimp"),
    ("Ciencias", r"ciencias(?! sociales)|biologia|fisica(?! y)|quimica|anatomia|steam"),
    ("Individuos y Sociedades", r"individuos|historia(?! del arte)|geografia|economia|empresa|gestion|politica|derecho"),
    ("Artes Visuales y Diseño", r"artes visuales|diseno|dibujo|fundamentos artisticos|historia del arte|programaci"),
    ("Artes Escénicas", r"musica|teatro|dramatizacion|conjunto instrumental"),
    ("Educación Física", r"educacion fisica|educacion personal, social y fisica|ajedrez"),
]


def departamento_de(asignatura: str) -> str:
    a = k(asignatura) + " "
    for dep, patron in DEPARTAMENTO:
        if re.search(patron, a):
            return dep
    return "Otros"


def grupo(g: str) -> str | None:
    m = re.fullmatch(r"(Y\d{1,2})[-\s]?([A-Z])(\.Comp)?", str(g or "").strip(), flags=re.I)
    return f"{m.group(1).upper()} {m.group(2).upper()}" if m else None


def partir_nombre(completo: str, correo: str) -> tuple[str, str, bool]:
    """(nombre, apellidos, seguro). El correo es nombre.apellidos (pegados, «ñ» como «nn»):
    los apellidos empiezan en la primera palabra desde la que el nombre completo, pegado,
    empieza por esa parte del correo."""
    def norm(x: str) -> str:
        return re.sub(r"[^a-z]", "", sin_acentos(x.replace("ñ", "nn").replace("Ñ", "nn")).lower())
    partes = completo.split()
    local = correo.split("@")[0]
    apellido = norm(local.split(".")[-1]) if "." in local else ""
    claves = [norm(p) for p in partes]
    if apellido:
        for i in range(1, len(partes)):
            pegado = "".join(claves[i:])
            if pegado.startswith(apellido) or (len(claves[i]) >= 3 and apellido.startswith(claves[i])):
                return " ".join(partes[:i]), " ".join(partes[i:]), True
    corte = 1 if len(partes) <= 3 else 2
    return " ".join(partes[:corte]), " ".join(partes[corte:]), False


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("alexia", type=Path, nargs="?", default=ORIGEN)
    ap.add_argument("--salida", type=Path, default=SALIDA)
    ap.add_argument("--revision", type=Path, default=REVISION,
                    help="revisión del colegio (hoja Departamentos; «Departamento correcto» manda)")
    args = ap.parse_args()

    import openpyxl
    import psycopg
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    with psycopg.connect(url_bd(False)) as conn, conn.cursor() as cur:
        cur.execute("""select distinct g.nombre, q.codigo from privado.respuesta r
                       join public.grupo g on g.id = r.grupo_id join public.cuestionario q on q.id = r.cuestionario_id
                       join public.curso c on c.id = q.curso_id where c.nombre = %s and q.colectivo = 'alumnos'""", (CURSO,))
        cuest_de = dict(cur.fetchall())
        cur.execute("""select distinct q.codigo, a.nombre from privado.respuesta_item i
                       join privado.respuesta r on r.id = i.respuesta_id join public.cuestionario q on q.id = r.cuestionario_id
                       join public.asignatura a on a.id = i.asignatura_id join public.curso c on c.id = q.curso_id
                       where c.nombre = %s""", (CURSO,))
        en_encuesta = defaultdict(set)
        for q, a in cur.fetchall():
            en_encuesta[q].add(k(a))

    # Comprobación: toda asignatura de la tabla existe en la encuesta de su cuestionario.
    for q, t in CORRESPONDENCIA.items():
        for destino in set(t.values()):
            if k(destino) not in en_encuesta[q]:
                raise SystemExit(f"{q}: {destino!r} no es una asignatura de la encuesta")

    libro = openpyxl.load_workbook(args.alexia, read_only=True, data_only=True)
    filas = list(libro["Detalle docente"].iter_rows(values_only=True))
    cab = [str(c or "").strip() for c in filas[0]]
    i = {c: cab.index(c) for c in ["Etapa", "Grupo", "Asignatura", "Docente", "Correo", "Profesor/a principal"]}

    asignacion: dict[tuple, dict] = {}
    sin_encuesta: Counter = Counter()
    docentes: dict[str, dict] = {}
    deps: dict[str, Counter] = defaultdict(Counter)
    imparte: dict[str, Counter] = defaultdict(Counter)
    for f in filas[1:]:
        correo = str(f[i["Correo"]] or "").strip().lower()
        nombre = str(f[i["Docente"]] or "").strip()
        if not correo:
            continue
        docentes.setdefault(correo, {"completo": nombre, "etapas": set()})["etapas"].add(f[i["Etapa"]])
        asig_alexia = str(f[i["Asignatura"]] or "").strip()
        if not k(asig_alexia).startswith(("tutoria", "refuerzo", "final")):
            deps[correo][departamento_de(asig_alexia)] += 1
            imparte[correo][k(asig_alexia) + " "] += 1
        g = grupo(f[i["Grupo"]])
        q = cuest_de.get(g or "")
        if q is None:
            continue
        destino = CORRESPONDENCIA[q].get(k(asig_alexia))
        if destino is None:
            sin_encuesta[(q, re.sub(r"-\s*Y\d{1,2}$", "", asig_alexia))] += 1
            continue
        clave_fila = (g, destino, correo)
        d = asignacion.setdefault(clave_fila, {"alexia": set(), "principal": False})
        d["alexia"].add(asig_alexia)
        d["principal"] |= str(f[i["Profesor/a principal"]]).strip().lower() == "sí"

    # Departamento propuesto: el de la mayoría de lo que imparte.
    propuestos = {}
    for correo, cnt in deps.items():
        total = sum(cnt.values())
        dep, n = cnt.most_common(1)[0] if cnt else ("Otros", 0)
        propuestos[correo] = (dep, "segura" if total and n / total >= 0.6 else "a revisar",
                              ", ".join(f"{d} ({m})" for d, m in cnt.most_common()))

    # Revisión del colegio: «Departamento correcto» manda sobre la propuesta; nombres oficiales.
    situacion = {str(r[1]).strip().lower(): str(r[2] or "Alta").strip()
                 for r in list(libro["Docentes"].iter_rows(values_only=True))[1:] if r[1]}
    finales: dict[str, tuple[str, str]] = {}
    revisar_ingles = []
    if args.revision.exists():
        rev = list(openpyxl.load_workbook(args.revision, read_only=True, data_only=True)["Departamentos"].iter_rows(values_only=True))
        cr = [str(c or "").strip() for c in rev[0]]
        for r in rev[1:]:
            d = dict(zip(cr, r))
            correo = str(d.get("Correo") or "").strip().lower()
            if not correo:
                continue
            correcto = str(d.get("Departamento correcto") or "").strip()
            dep = correcto or str(d.get("Departamento propuesto") or OTROS).strip()
            finales[correo] = (RENOMBRAR.get(dep, RENOMBRAR.get(dep.lower(), dep)), "colegio" if correcto else "propuesta confirmada")
    for correo, (dep, origen) in list(finales.items()):
        if dep != INGLES_Y_ADQUISICION:
            continue
        ing = sum(n for a, n in imparte[correo].items() if re.search(INGLESAS, a))
        adq = sum(n for a, n in imparte[correo].items() if re.search(ADQUISICION, a))
        nuevo = "Lengua y Literatura Inglesas" if ing > adq else "Adquisición de Lenguas"
        if ing and adq or not (ing or adq):
            revisar_ingles.append((correo, ing, adq, nuevo))
            finales[correo] = (nuevo, "a revisar: imparte de los dos" if ing and adq else "a revisar: no encaja en ninguno")
        else:
            finales[correo] = (nuevo, origen + " (reparto de Inglés y Adquisición)")
    for correo, (dep, _) in finales.items():
        if dep not in OFICIALES and dep != OTROS:
            raise SystemExit(f"Departamento no oficial para {correo}: {dep!r}")
    if finales:
        propuestos = {c: (finales[c][0], finales[c][1], propuestos.get(c, (None, None, ""))[2]) if c in finales
                      else propuestos.get(c, (OTROS, "a revisar", "")) for c in docentes}

    salida = openpyxl.Workbook()
    negrita = Font(bold=True, color="FFFFFF")
    azul = PatternFill("solid", fgColor="1F3A5F")

    def cabecera(ws, columnas, anchos):
        ws.append(columnas)
        for c, a in zip(ws[1], anchos):
            c.font, c.fill = negrita, azul
            ws.column_dimensions[c.column_letter].width = a
        ws.freeze_panes = "A2"

    hd = salida.active
    hd.title = "Departamentos"
    cabecera(hd, ["Correo", "Apellidos", "Nombre", "Departamento", "Propuesta", "Lo que imparte (clases por departamento)", "Etapas", "Nombre partido", "Situación"],
             [34, 26, 22, 30, 30, 70, 40, 14, 10])
    for correo in sorted(docentes, key=lambda c: partir_nombre(docentes[c]["completo"], c)[1]):
        nombre, apellidos, seguro = partir_nombre(docentes[correo]["completo"], correo)
        dep, confianza, base = propuestos.get(correo, ("Otros", "a revisar", ""))
        hd.append([correo, apellidos, nombre, dep, confianza, base, ", ".join(sorted(docentes[correo]["etapas"])),
                   "sí" if seguro else "revisar", situacion.get(correo, "Alta")])
    lista = DataValidation(type="list", formula1='"' + ",".join(OFICIALES + [OTROS]) + '"')
    hd.add_data_validation(lista)
    lista.add(f"D2:D{hd.max_row}")

    ha = salida.create_sheet("Asignación", 0)
    cabecera(ha, ["Curso", "Grupo", "Asignatura", "Apellidos", "Nombre", "Correo", "Departamento"], [10, 10, 60, 26, 22, 34, 30])
    for (g, destino, correo) in sorted(asignacion, key=lambda x: (int(x[0][1:].split()[0]), x[0], x[1], x[2])):
        nombre, apellidos, _ = partir_nombre(docentes[correo]["completo"], correo)
        ha.append([CURSO, g, destino, apellidos, nombre, correo, propuestos.get(correo, ("Otros",))[0]])

    hc = salida.create_sheet("Codocencia")
    cabecera(hc, ["Grupo", "Asignatura de la encuesta", "Docentes", "Materias de Alexia"], [10, 60, 70, 70])
    por_comb = defaultdict(list)
    for (g, destino, correo), d in asignacion.items():
        por_comb[(g, destino)].append((docentes[correo]["completo"] + (" (principal)" if d["principal"] else ""), d["alexia"]))
    for (g, destino), lista_d in sorted(por_comb.items()):
        if len(lista_d) > 1:
            hc.append([g, destino, "; ".join(sorted(n for n, _ in lista_d)),
                       "; ".join(sorted({a for _, s in lista_d for a in s}))])

    hs = salida.create_sheet("Sin encuesta")
    cabecera(hs, ["Cuestionario", "Materia de Alexia", "Filas", "Motivo"], [14, 50, 8, 80])
    for (q, a), n in sorted(sin_encuesta.items()):
        hs.append([q, a, n, "Los alumnos no la valoran en su encuesta"])
    for q, motivo in NO_ATRIBUIBLES.items():
        hs.append(["", f"(pregunta de la encuesta) {q}", "", motivo])

    args.salida.parent.mkdir(parents=True, exist_ok=True)
    salida.save(args.salida)

    # Excel corto para el colegio: docentes del antiguo «Inglés y Adquisición» por revisar.
    if revisar_ingles:
        rv = openpyxl.Workbook()
        hr = rv.active
        hr.title = "Revisar"
        cabecera(hr, ["Docente", "Lo que imparte", "Departamento propuesto", "Departamento correcto"], [32, 80, 30, 30])
        for correo, ing, adq, nuevo in sorted(revisar_ingles, key=lambda x: partir_nombre(docentes[x[0]]["completo"], x[0])[1]):
            materias = ", ".join(f"{a.strip().capitalize()} ({n})" for a, n in imparte[correo].most_common())
            hr.append([docentes[correo]["completo"], materias, nuevo, None])
        for c in hr["B"]:
            c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
        lista2 = DataValidation(type="list", formula1='"Lengua y Literatura Inglesas,Adquisición de Lenguas,Otros"')
        hr.add_data_validation(lista2)
        lista2.add(f"D2:D{hr.max_row}")
        hr.append([])
        hr.append(["Imparten materias de Lengua y Literatura Inglesas y de Adquisición de Lenguas (o de ninguna de las dos). "
                   "Se han cargado en el departamento donde dan más clases. Si no es el correcto, elegidlo en la última columna."])
        rv.save(REVISAR_INGLES)
        print(f"Para revisar (Inglés y Adquisición): {len(revisar_ingles)} docentes → {REVISAR_INGLES.relative_to(RAIZ)}")

    revisar_dep = sum(1 for c in docentes if propuestos.get(c, ("", "a revisar"))[1] == "a revisar")
    revisar_nom = sum(1 for c in docentes if not partir_nombre(docentes[c]["completo"], c)[2])
    con_asig = {c for _, _, c in asignacion}
    print(f"Docentes en Alexia: {len(docentes)}; con alguna asignatura de las encuestas: {len(con_asig)}")
    print(f"Filas de asignación: {len(asignacion)}; combinaciones con codocencia: {sum(1 for v in por_comb.values() if len(v) > 1)} de {len(por_comb)}")
    print(f"Departamento a revisar: {revisar_dep}; nombre y apellidos a revisar: {revisar_nom}")
    print(f"Salida: {args.salida.relative_to(RAIZ) if args.salida.is_relative_to(RAIZ) else args.salida}")


if __name__ == "__main__":
    main()
